import io
import re

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pypdf import PdfReader

from app import config, db, impresora, inventario as inv
from app.main import app


def cliente():
    return TestClient(app)


def nuevo_producto(c, nombre="Molde vela corazón", cantidad="6", costo="52.542", **extra):
    datos = {"nombre_interno": nombre, "marca": "Super K", "categoria_id": "1",
             "proveedor": "Super K", "nombre_proveedor": "MOLDE VELA CORAZON P J27592",
             "fecha": "2026-06-05", "nro_factura": "001-001-0133014", "cantidad": cantidad,
             "costo_total": costo, "margen_pct": "40", "precio_venta": ""}
    datos.update(extra)
    return c.post("/productos/nuevo", data=datos)


def textos_pdf(contenido: bytes) -> str:
    return "\n".join(pagina.extract_text() for pagina in PdfReader(io.BytesIO(contenido)).pages)


def test_flujo_completo_web():
    with cliente() as c:
        r = c.get("/")
        assert "No hay avisos" in r.text
        r = c.get("/productos/nuevo")
        assert "P-00001" in r.text and "SKU" not in r.text.split("<form")[1]

        r = nuevo_producto(c)
        assert r.status_code == 200 and "Producto guardado" in r.text
        assert "P-00001" in r.text
        # la tabla ya no muestra SKU ni código de barra, y el proveedor es el de la última compra
        tabla = r.text.split('id="tabla-productos"')[1]
        assert "SKU" not in tabla and "/barra/" not in tabla and "Última compra en" in tabla
        assert 'value="12260"' in r.text  # 8.757 + 40 %
        assert "Imprimir lista" not in r.text and "Cambiar precios por porcentaje" not in r.text

        r = c.get("/compras")
        assert "001-001-0133014" in r.text and "₲ 8.757" in r.text and "₲ 52.542" in r.text
        assert "Agregar compra" not in r.text and "/compras/panel" in r.text
        assert c.get("/compras/nueva").status_code in (404, 405)

        assert c.get("/api/stock").json()["stock"]["1"] == 6
        assert c.get("/barra/20022001.svg").headers["content-type"].startswith("image/svg")

        r = c.post("/productos/1/editar", data={
            "nombre_interno": "Molde vela corazón P", "categoria_id": "1", "proveedor": "Super K",
            "margen_pct": "50", "precio_venta": "13.136", "activo": "1"})
        assert "Cambios guardados" in r.text and 'value="13136"' in r.text


def test_precio_por_porcentaje_o_a_mano_en_la_tabla():
    with cliente() as c:
        nuevo_producto(c, cantidad="1", costo="10000")
        r = c.get("/productos")
        assert 'name="precio_venta"' in r.text and 'name="margen_pct"' in r.text
        # por porcentaje: se calcula el precio
        r = c.post("/productos/1/precio", data={"margen_pct": "100", "precio_venta": "20000", "origen": "margen"})
        assert "precio ₲ 20.000 (ganancia 100%)" in r.text
        # a mano: se calcula el porcentaje
        r = c.post("/productos/1/precio", data={"margen_pct": "100", "precio_venta": "13.500", "origen": "precio"})
        assert "precio ₲ 13.500 (ganancia 35%)" in r.text
        assert 'value="35"' in r.text and 'value="13500"' in r.text and 'value="13500"' in r.text
        # sin decir cuál se tocó: vale el precio si cambió
        r = c.post("/productos/1/precio", data={"margen_pct": "35", "precio_venta": "15.550"})
        assert "precio ₲ 15.550 (ganancia 55,5%)" in r.text
        # la ruta vieja por % sigue funcionando
        r = c.post("/productos/1/margen", data={"margen_pct": "40"})
        assert "precio ₲ 14.000 (ganancia 40%)" in r.text
        r = c.post("/productos/1/precio", data={"precio_venta": "0", "origen": "precio"})
        assert "precio de venta válido" in r.text
        r = c.post("/productos/1/precio", data={"margen_pct": "abc", "origen": "margen"})
        assert "porcentaje válido" in r.text


def test_venta_con_carrito_y_cliente_nuevo():
    with cliente() as c:
        nuevo_producto(c, "Molde A", cantidad="5", costo="5000", precio_venta="3.000")
        nuevo_producto(c, "Molde B", cantidad="2", costo="2000", precio_venta="2.500")
        r = c.get("/vender")
        assert 'id="datos-productos"' in r.text and "Molde A" in r.text

        nuevo = c.post("/api/clientes", json={"nombre": "Kari Personalizados", "telefono": "0981 1",
                                              "documento": "80012345-6"}).json()
        assert nuevo["texto"] == "Kari Personalizados" and "RUC/CI 80012345-6" in nuevo["detalle"]
        assert c.post("/api/clientes", json={"nombre": " "}).status_code == 400

        r = c.post("/vender", data={"producto_id": ["1", "2"], "cantidad": ["2", "1"],
                                    "precio_unit": "1", "cliente_id": str(nuevo["id"]),
                                    "medio_pago": "Efectivo", "fecha": "2026-06-26"})
        assert "Venta guardada: 2 producto(s) por ₲ 8.500" in r.text
        assert "Kari Personalizados" in r.text
        stock = c.get("/api/stock").json()["stock"]
        assert stock["1"] == 3 and stock["2"] == 1

        r = c.post("/vender", data={"producto_id": ["1", "2"], "cantidad": ["1", "5"]})
        assert "No alcanza el stock de «Molde B»" in r.text
        assert c.get("/api/stock").json()["stock"]["1"] == 3

        libro = load_workbook(io.BytesIO(c.get("/ventas/excel").content))
        assert libro.active.max_row == 3


def test_panel_de_compras_con_filtros():
    with cliente() as c:
        nuevo_producto(c, "Molde A", fecha="2026-05-10", proveedor="Pacific", costo="1000")
        nuevo_producto(c, "Molde B", fecha="2026-06-10", proveedor="Super K", costo="3000")
        r = c.get("/compras/panel")
        assert "Gasto por proveedor" in r.text and "₲ 4.000" in r.text
        r = c.get("/compras/panel?desde=2026-06-01")
        assert "Molde B" in r.text and "<b>Molde A</b>" not in r.text and "₲ 3.000" in r.text
        r = c.get("/compras/panel?producto_id=1")
        assert "<b>Molde A</b>" in r.text and "<b>Molde B</b>" not in r.text
        libro = load_workbook(io.BytesIO(c.get("/compras/excel?producto_id=2").content))
        assert libro.active.max_row == 2 and libro.active["E2"].value == "Molde B"


def test_etiquetas_tamano_pdf_e_impresora(monkeypatch):
    with cliente() as c:
        nuevo_producto(c, "Molde A", cantidad="3", costo="3000")
        r = c.get("/etiquetas")
        assert "Tamaño de la etiqueta" in r.text and "Cómo conectar la impresora" in r.text
        assert "🖨️ Imprimir</button>" not in r.text  # sin impresora configurada solo hay PDF

        r = c.post("/etiquetas/tamano", data={"tamano": "40x40"})
        assert "40 × 40 mm" in r.text and 'value="40x40" checked' in r.text
        assert 'viewBox="0 0 40 40"' in c.get("/etiquetas/vista.svg?producto_id=1").text

        r = c.get("/etiquetas/factura/1.pdf")
        lector = PdfReader(io.BytesIO(r.content))
        assert len(lector.pages) == 3  # una etiqueta por unidad, una por página del rollo
        assert abs(float(lector.pages[0].mediabox.width) - 40 * 72 / 25.4) < 0.5

        r = c.post("/etiquetas/producto", data={"producto_id": "1", "cantidad": "25", "accion": "a4"})
        assert len(PdfReader(io.BytesIO(r.content)).pages) == 2  # 21 por hoja A4
        assert textos_pdf(r.content).count("20022001") == 25 and "Gs. 1.400" in textos_pdf(r.content)

        r = c.post("/etiquetas/producto", data={"producto_id": "", "cantidad": "2", "accion": "pdf"})
        assert "Elegí el producto" in r.text

        enviados = []
        monkeypatch.setattr(config, "IMPRESORA_NOMBRE", "Xprinter XP-365B")
        monkeypatch.setattr(impresora, "enviar", lambda datos: enviados.append(datos))
        r = c.get("/etiquetas")
        assert "Xprinter XP-365B" in r.text and "🧪 Imprimir etiqueta de prueba" in r.text
        r = c.post("/etiquetas/factura/1/imprimir")
        assert "Se mandaron 3 etiqueta(s) de 40 × 40 mm" in r.text
        r = c.post("/etiquetas/producto", data={"producto_id": "1", "cantidad": "5", "accion": "imprimir"})
        assert "Se mandaron 5 etiqueta(s)" in r.text
        assert b"SIZE 40 mm,40 mm" in enviados[0] and b"PRINT 1,5" in enviados[1]


def test_grafico_de_costo_del_producto_con_filtros():
    with cliente() as c:
        nuevo_producto(c, "Molde A", fecha="2026-03-10", proveedor="Pacific", cantidad="10", costo="10000")
        for fecha, prov, costo in (("2026-04-10", "Pacific", "12000"), ("2026-05-10", "Super K", "9000"),
                                   ("2026-06-10", "Pacific", "15000")):
            with db.conexion() as con:
                inv.registrar_compra(con, fecha, "F", prov, [{"producto_id": 1, "cantidad": 10, "costo_total": costo}])
        r = c.get("/compras/panel?producto_id=1")
        assert "Costo por unidad" in r.text
        puntos = re.search(r'<polyline class="linea" points="([^"]+)"', r.text).group(1).split()
        assert len(puntos) == 4
        assert "▲" in r.text and "+50,0%" in r.text  # de 1.000 a 1.500
        pacific = re.search(r'<option value="(\d+)"[^>]*>Pacific<', r.text).group(1)
        r = c.get(f"/compras/panel?producto_id=1&proveedor_id={pacific}&desde=2026-04-01")
        assert len(re.search(r'points="([^"]+)"', r.text).group(1).split()) == 2
        r = c.get("/compras/panel?producto_id=1&desde=2026-06-01")
        assert "Hacen falta al menos 2 compras" in r.text


def test_unir_desde_la_pagina():
    with cliente() as c:
        nuevo_producto(c, "Pistola de calor", cantidad="2", costo="80000")
        nuevo_producto(c, "PISTOLA CALOR XQ", cantidad="3", costo="120000")
        r = c.get("/productos/2/editar")
        assert "¿Este producto ya existía con otro nombre?" in r.text
        r = c.post("/productos/2/unir", data={"destino_id": "1"})
        assert "la compra pasó a «Pistola de calor». Ahora hay 5" in r.text
        assert c.get("/productos/2/editar").url.path == "/productos"
        r = c.post("/productos/1/unir", data={"destino_id": ""})
        assert "Elegí de la lista" in r.text


def test_menu_y_compras_sin_botones_quitados():
    with cliente() as c:
        r = c.get("/compras")
        assert "Nuevo producto</span>" not in r.text and "Imprimir</button>" not in r.text
        assert "✨ Nuevo producto" in c.get("/productos").text


def test_notificaciones_con_badge():
    with cliente() as c:
        nuevo_producto(c, "Molde A", cantidad="1", costo="1000")
        r = c.get("/")
        assert "Queda poco stock" in r.text
        assert re.search(r'class="globo"[^>]*>1<', r.text)


def test_pide_clave(monkeypatch):
    monkeypatch.setattr(config, "CLAVE_ACCESO", "1234")
    with cliente() as c:
        r = c.get("/productos")
        assert r.url.path == "/login"
        r = c.post("/login", data={"clave": "mal"})
        assert "no es correcta" in r.text
        r = c.post("/login", data={"clave": "1234"})
        assert r.url.path == "/"
        assert c.get("/productos").url.path == "/productos"


def test_editar_sin_duplicados_y_menu_fijo():
    with cliente() as c:
        nuevo_producto(c, "Molde A", cantidad="1", costo="1000")
        r = c.get("/productos/1/editar")
        assert "<title>Editar producto</title>" in r.text
        assert r.text.count('id="buscador-unir"') == 1
        assert '<header class="barra-menu no-imprimir">' in r.text



def test_fotos_con_protagonismo_en_stock_y_vender():
    import io as _io
    from PIL import Image
    foto = _io.BytesIO()
    Image.new("RGB", (300, 300), "orange").save(foto, format="JPEG")
    with cliente() as c:
        nuevo_producto(c, "Angelito tipo 1", cantidad="3", costo="3000")
        c.post("/productos/1/editar", data={"nombre_interno": "Angelito tipo 1", "categoria_id": "1",
                                            "margen_pct": "40", "precio_venta": "1400", "activo": "1"},
               files={"imagen": ("a.jpg", foto.getvalue(), "image/jpeg")})
        r = c.get("/productos")
        assert 'class="miniatura ampliable"' in r.text and 'alt="Angelito tipo 1"' in r.text
        r = c.get("/vender")
        assert '"imagen": "/archivo/productos/' in r.text  # la foto aparece en el buscador
        c.post("/vender", data={"producto_id": ["1"], "cantidad": ["1"]})
        assert 'class="foto-chica ampliable"' in c.get("/vender").text  # y en las últimas ventas
        css = c.get("/static/estilo.css").text
        assert "[hidden] { display: none !important; }" in css  # el visor oculto no tapa la página
