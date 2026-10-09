"""Exportación a Excel (.xlsx)."""
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app import inventario as inv

FORMATO_GS = '#,##0'


def _libro(titulo: str, encabezados: list[str], filas: list[list], columnas_dinero=()) -> bytes:
    libro = Workbook()
    hoja = libro.active
    hoja.title = titulo
    hoja.append(encabezados)
    for celda in hoja[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="2F5D50")
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for fila in filas:
        hoja.append(fila)
    for indice in columnas_dinero:
        letra = get_column_letter(indice)
        for celda in hoja[letra][1:]:
            celda.number_format = FORMATO_GS
    for i, encabezado in enumerate(encabezados, 1):
        largo = max([len(str(encabezado))] + [len(str(f[i - 1] or "")) for f in filas])
        hoja.column_dimensions[get_column_letter(i)].width = min(max(largo + 2, 10), 60)
    hoja.freeze_panes = "A2"
    hoja.auto_filter.ref = hoja.dimensions
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


def productos(con) -> bytes:
    encabezados = ["Código", "Nombre interno", "Nombre en la factura",
                   "Marca", "Categoría", "Última compra en", "Costo última compra", "Margen %",
                   "Precio venta", "En stock", "Activo"]
    filas = [[p["codigo"], p["nombre_interno"], p["nombre_proveedor"],
              p["marca"], p["categoria"], p["proveedor"], round(p["costo_ultimo"]),
              p["margen_pct"], p["precio_venta"], p["stock"], "Sí" if p["activo"] else "No"]
             for p in inv.listar_productos(con)]
    return _libro("Productos", encabezados, filas, columnas_dinero=(7, 9))


def compras(con, filtros: dict | None = None) -> bytes:
    encabezados = ["Fecha", "Código producto", "N° factura", "Proveedor", "Producto",
                   "Cantidad", "Precio unitario", "Precio total", "Queda en stock", "Origen"]
    filas = [[inv.fecha_txt(c["fecha"]), c["codigo"], c["nro_factura"], c["proveedor"],
              c["nombre_interno"], c["cantidad"], round(c["costo_unitario"]),
              round(c["costo_total"]), c["cantidad_restante"], c["origen"]]
             for c in inv.listar_compras(con, **(filtros or {}))]
    return _libro("Compras", encabezados, filas, columnas_dinero=(7, 8))


def ventas(con) -> bytes:
    encabezados = ["Fecha", "Cliente", "Código", "Producto", "Cantidad", "Precio unit.",
                   "Total", "Costo (FIFO)", "Ganancia", "Medio de pago"]
    filas = [[inv.fecha_txt(v["fecha"]), v["cliente"], v["codigo"], v["nombre_interno"],
              v["cantidad"], round(v["precio_unit"]), round(v["total"]), round(v["costo_fifo"]),
              round(v["ganancia"]), v["medio_pago"]]
             for v in inv.listar_ventas(con)]
    return _libro("Ventas", encabezados, filas, columnas_dinero=(6, 7, 8, 9))
