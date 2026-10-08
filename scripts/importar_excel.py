"""Carga inicial desde el Excel anterior (hojas Producto, Compras y Ventas).

Uso:  python scripts/importar_excel.py "Datos.xlsx"  [--forzar]

- Producto: cada producto con su SKU, código de barra y precio; el "Stock disponible"
  entra como un lote "Stock inicial" al "Costo Promedio".
- Compras: si el SKU no está en la hoja Producto se crea el producto y la compra suma stock.
  Si ya estaba, la compra se guarda como historial (no suma, porque ya está en el stock inicial).
- Ventas: se guardan como historial. Solo descuentan stock las de productos creados desde
  la hoja Compras (los de la hoja Producto ya tienen las ventas descontadas).
"""
import argparse
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from openpyxl import load_workbook  # noqa: E402

from app import db, inventario as inv  # noqa: E402


def filas_de(hoja):
    filas = hoja.iter_rows(values_only=True)
    encabezados = [str(c).strip() if c is not None else "" for c in next(filas)]
    for fila in filas:
        yield dict(zip(encabezados, fila))


def numero(valor, por_defecto=0.0):
    if isinstance(valor, str) and valor.startswith("#"):
        return por_defecto  # celdas con error del Excel (#VALUE!)
    n = inv.leer_numero(valor)
    return por_defecto if n is None else n


def importar(ruta: str, forzar: bool = False) -> list[str]:
    avisos: list[str] = []
    db.inicializar()
    libro = load_workbook(ruta, data_only=True)
    with db.conexion() as con:
        if con.execute("SELECT COUNT(*) FROM productos").fetchone()[0] and not forzar:
            raise SystemExit("La base ya tiene productos. Usá --forzar si igual querés importar.")

        hoja_compras = list(filas_de(libro["Compras"])) if "Compras" in libro.sheetnames else []
        fechas = [inv.leer_fecha(f["Fecha"]) for f in hoja_compras if f.get("Fecha")]
        fecha_inicial = min(fechas) if fechas else None

        # ---------------- Producto
        desde_hoja_producto: set[str] = set()
        por_proveedor: dict[str, list] = {}
        for f in filas_de(libro["Producto"]):
            sku = (f.get("SKU") or "").strip()
            if not sku or not f.get("Producto"):
                continue
            if con.execute("SELECT 1 FROM productos WHERE sku = ?", (sku,)).fetchone():
                avisos.append(f"Producto {sku}: ya existía, no se volvió a cargar.")
                continue
            margen = numero(f.get("Margen %"), None)
            proveedor = (f.get("Proveedor") or "").strip()
            creado = inv.crear_producto(con, {
                "sku": sku,
                "codigo_barra": str(f.get("Codigo de barra") or "").strip(),
                "nombre_interno": str(f["Producto"]).strip(),
                "nombre_proveedor": str(f["Producto"]).strip(),
                "marca": (f.get("Marca") or "").strip(),
                "categoria": f.get("Categoría") or "Varios",
                "proveedor": proveedor,
                "costo_unitario": numero(f.get("Costo Promedio")),
                "margen_pct": round(margen * 100, 1) if margen is not None else None,
                "precio_venta": numero(f.get("Precio Venta"), None),
                "activo": str(f.get("Activo") or "Si").strip().lower().startswith("s"),
            })
            desde_hoja_producto.add(sku)
            disponible = numero(f.get("Stock disponible"))
            if disponible > 0:
                por_proveedor.setdefault(proveedor, []).append({
                    "producto_id": creado["id"], "cantidad": disponible,
                    "costo_total": disponible * numero(f.get("Costo Promedio")),
                })
            elif disponible < 0:
                avisos.append(f"Producto {sku}: el stock disponible era negativo ({disponible}).")
        for proveedor, items in por_proveedor.items():
            inv.registrar_compra(con, fecha_inicial, "STOCK INICIAL", proveedor, items,
                                 origen="excel")

        # ---------------- Compras
        for f in hoja_compras:
            sku = (f.get("SKU") or "").strip()
            if not sku:
                continue
            cantidad = numero(f.get("Cantidad"))
            costo_total = numero(f.get("Costo Total")) or cantidad * numero(f.get("Costo Unitario"))
            if cantidad <= 0:
                avisos.append(f"Compra {sku}: cantidad vacía, se saltó.")
                continue
            fila = con.execute("SELECT id FROM productos WHERE sku = ?", (sku,)).fetchone()
            if fila and sku in desde_hoja_producto:
                restante = 0  # ya está contada en el stock inicial
                producto_id = fila["id"]
            else:
                if not fila:
                    producto_id = inv.crear_producto(con, {
                        "sku": sku, "nombre_interno": str(f.get("Producto") or sku).strip(),
                        "categoria": "Varios", "proveedor": f.get("Proveedor") or "",
                        "costo_unitario": costo_total / cantidad, "revisar": True,
                    })["id"]
                    avisos.append(f"Compra {sku}: el producto no estaba en la hoja Producto; "
                                  "se creó en categoría Varios (marcado para revisar).")
                else:
                    producto_id = fila["id"]
                restante = cantidad
            inv.registrar_compra(con, f.get("Fecha"), str(f.get("Factura") or ""),
                                 f.get("Proveedor") or "",
                                 [{"producto_id": producto_id, "cantidad": cantidad,
                                   "costo_total": costo_total, "cantidad_restante": restante}],
                                 origen="excel")

        # ---------------- Ventas
        if "Ventas" in libro.sheetnames:
            for f in filas_de(libro["Ventas"]):
                sku = (f.get("SKU") or "").strip()
                if not sku:
                    continue
                fila = con.execute("SELECT * FROM productos WHERE sku = ?", (sku,)).fetchone()
                if not fila:
                    avisos.append(f"Venta {sku} ({f.get('Producto')}): el producto no existe, "
                                  "no se cargó.")
                    continue
                cantidad = numero(f.get("Cantidad"))
                precio = numero(f.get("Precio Unit."))
                nombre_cliente = str(f.get("Cliente") or "").strip()
                cliente_id = inv.crear_cliente(con, nombre_cliente)["id"] if nombre_cliente else None
                medio = (f.get("Medio Pago") or "").strip().capitalize()
                if sku in desde_hoja_producto:
                    # Ya está descontada del stock disponible: solo se guarda como historial.
                    venta_id = con.execute(
                        "INSERT INTO ventas (fecha, cliente_id, medio_pago) VALUES (?, ?, ?)",
                        (inv.leer_fecha(f.get("Fecha")), cliente_id, medio),
                    ).lastrowid
                    costo = cantidad * inv.costo_referencia(con, fila["id"])
                    con.execute(
                        """INSERT INTO venta_items (venta_id, producto_id, cantidad, precio_unit,
                               costo_fifo) VALUES (?, ?, ?, ?, ?)""",
                        (venta_id, fila["id"], cantidad, precio, costo),
                    )
                else:
                    try:
                        r = inv.registrar_venta(con, [{"producto_id": fila["id"], "cantidad": cantidad}],
                                                cliente_id, medio, f.get("Fecha"))
                        # Se respeta el precio al que se vendió en el Excel.
                        con.execute("UPDATE venta_items SET precio_unit = ? WHERE venta_id = ?",
                                    (precio, r["venta_id"]))
                    except inv.ErrorNegocio as exc:
                        avisos.append(f"Venta {sku}: {exc}")

        total = con.execute("SELECT COUNT(*) FROM productos").fetchone()[0]
        unidades = con.execute("SELECT COALESCE(SUM(cantidad_restante), 0) FROM lotes").fetchone()[0]
        avisos.insert(0, f"Listo: {total} productos, {inv.cantidad_txt(unidades)} unidades en stock.")
    return avisos


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__,
                                     formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("excel")
    parser.add_argument("--forzar", action="store_true")
    args = parser.parse_args()
    for linea in importar(args.excel, args.forzar):
        print("•", linea)
